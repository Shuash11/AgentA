import hashlib
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "tools" / "revise_agenta_traceability_matrix.py"
SOURCE = Path(
    r"C:\Users\joashua\Downloads\AgentA_FR_Design_Pattern_Traceability_Matrix (1).docx"
)
OUTPUT = Path(
    r"C:\Users\joashua\Downloads\AgentA_FR_Design_Pattern_Traceability_Matrix_REVISED.docx"
)

EXPECTED_PARAGRAPHS = [
    "AgentA: An AI-Based Learning Assistant",
    "Functional Requirement-Design Pattern Traceability Matrix",
    (
        "This matrix links AgentA's main functional requirements to the software "
        "design patterns used in the system. For each requirement, it states what "
        "the pattern does, where it is used, and why it suits the required behavior. "
        "Technical names appear only when they are needed to show traceability."
    ),
    "Summary",
    (
        "AgentA uses the Facade pattern for account, profile, recovery, and security "
        "functions. This lets each screen call a service instead of managing every "
        "step itself. The Observer pattern updates the chat interface and theme when "
        "their state changes, while the Strategy pattern selects an online or offline "
        "response method based on connection status. The Repository, Chain of "
        "Responsibility, Command, and State patterns support chat-session management, "
        "password recovery, account deletion, and navigation."
    ),
]

EXPECTED_ROWS = [
    [
        "Functional Requirement",
        "Design Pattern",
        "Purpose of the Pattern",
        "Application in the System",
        "Justification",
    ],
    [
        "User Registration & Recovery Code Generation",
        "Facade Pattern",
        "Presents registration through one service that coordinates input checks, account creation, and recovery code generation.",
        "The registration screen uses AuthService to manage account validation and creation instead of handling each step itself.",
        "Registration has several related steps. Placing them in AuthService avoids repeating the same work in the screen and gives users a consistent registration process.",
    ],
    [
        "User Login & Session Handling",
        "Facade Pattern + Singleton Pattern",
        "The Facade pattern gives the screen one point of access to authentication. The Singleton pattern keeps one shared authentication state while the application is running.",
        "AuthService checks the login details and establishes the current user session before the Dashboard opens.",
        "A user should receive one clear login result. Centralized authentication and one shared session state reduce the risk of conflicting account information.",
    ],
    [
        "Password Recovery via Recovery Code",
        "Chain of Responsibility Pattern",
        "Arranges password recovery as a sequence in which one stage must be completed before the next can begin.",
        "RecoveryCodeService checks the recovery code first. A user can set a new password only after the code is accepted.",
        "The sequence prevents a password change before identity verification. If verification fails, the process stops and the user receives an error.",
    ],
    [
        "Chat Session Management (Dashboard)",
        "Repository Pattern",
        "Centralizes the creation, retrieval, and deletion of chat-session records.",
        "The Dashboard manages the user's chat sessions and provides access to previous conversations.",
        "A central collection keeps conversation history organized and removes data-management work from the interface.",
    ],
    [
        "Conversational Chat Interface (Message Display)",
        "Observer Pattern",
        "Updates parts of the interface automatically when the conversation state changes.",
        "The message list, typing indicator, and conversation position change as messages are sent and received.",
        "The interface must show each exchange as it happens. State-based updates keep the conversation current without a manual refresh.",
    ],
    [
        "Dual-Mode Response Engine (Online/Offline Routing)",
        "Strategy Pattern",
        "Defines alternative response methods and selects one without changing how the user sends a message.",
        "The response engine uses the online AI service when an internet connection is available. When the device is offline, it uses a local keyword-based response.",
        "Both modes answer the user's message, but they rely on different resources. Choosing the method at runtime allows the chat to remain available when the connection changes.",
    ],
    [
        "Profile Viewing and Editing",
        "Facade Pattern",
        "Places profile validation, saving, and confirmation behind one service.",
        "ProfileService manages changes submitted through the Edit Profile screen and confirms when they have been saved.",
        "The screen can focus on user input while ProfileService handles the rules for checking and saving profile information.",
    ],
    [
        "Security Settings (View Recovery Code / Delete Account)",
        "Facade Pattern + Command Pattern",
        "The Facade pattern gives the screen controlled access to security functions. The Command pattern treats account deletion as a separate action that must be confirmed before it runs.",
        "SecurityService allows users to view the recovery code or delete an account. Deletion starts only after the required confirmation steps are completed.",
        "Security functions need one controlled access point. Treating deletion as a separate, confirmable action reduces the chance of accidental data loss.",
    ],
    [
        "Theme Management (Dark/Light Mode Toggle)",
        "Observer Pattern",
        "Notifies relevant screens when a shared display preference changes.",
        "ThemeNotifier stores the light or dark theme choice and applies it across the application.",
        "The theme affects several screens. A shared notification mechanism applies the same choice everywhere as soon as the user changes it.",
    ],
    [
        "Bottom Navigation Bar (Tab Switching)",
        "State Pattern",
        "Changes navigation behavior according to the tab that is currently selected.",
        "The bottom navigation bar tracks the active tab, opens a newly selected section, and does nothing when the current tab is selected again.",
        "The result of a tap depends on the active tab. Tracking that state prevents unnecessary navigation and makes the behavior predictable.",
    ],
]


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class AgentATraceabilityMatrixRevisionTests(unittest.TestCase):
    def test_creates_readable_formal_copy_without_changing_source(self):
        self.assertTrue(SOURCE.exists(), f"Source document is missing: {SOURCE}")
        source_hash = sha256(SOURCE)

        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "revised.docx"
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(SOURCE), str(output)],
                capture_output=True,
                text=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
            self.assertTrue(output.exists())

            source_document = Document(SOURCE)
            document = Document(output)
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            rows = [
                [cell.text for cell in row.cells]
                for row in document.tables[0].rows
            ]

            self.assertEqual(paragraphs, EXPECTED_PARAGRAPHS)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(len(rows), 11)
            self.assertTrue(all(len(row) == 5 for row in rows))
            self.assertEqual(rows, EXPECTED_ROWS)
            self.assertEqual(
                [p.style.name if p.style else None for p in source_document.paragraphs],
                [p.style.name if p.style else None for p in document.paragraphs],
            )
            self.assertEqual(
                source_document.tables[0].style.name
                if source_document.tables[0].style
                else None,
                document.tables[0].style.name if document.tables[0].style else None,
            )
            self.assertEqual(
                [
                    cell.paragraphs[0].style.name
                    if cell.paragraphs[0].style
                    else None
                    for row in source_document.tables[0].rows
                    for cell in row.cells
                ],
                [
                    cell.paragraphs[0].style.name
                    if cell.paragraphs[0].style
                    else None
                    for row in document.tables[0].rows
                    for cell in row.cells
                ],
            )

            combined_text = "\n".join(paragraphs + [cell for row in rows for cell in row])
            for prohibited in (
                ".dart",
                "()",
                "ChangeNotifierProvider",
                "PageRouteBuilder",
                "_generateReply",
                "1.2 seconds",
                "1-second delay",
                "\ufffd",
                "\u2013",
                "\u2014",
            ):
                self.assertNotIn(prohibited, combined_text)

            for required in (
                "AuthService",
                "RecoveryCodeService",
                "ProfileService",
                "SecurityService",
                "ThemeNotifier",
            ):
                self.assertIn(required, combined_text)

            with ZipFile(output) as archive:
                self.assertIsNone(archive.testzip())

        self.assertEqual(sha256(SOURCE), source_hash)

    def test_download_copy_contains_the_approved_revision(self):
        self.assertTrue(OUTPUT.exists(), f"Revised document is missing: {OUTPUT}")

        document = Document(OUTPUT)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        rows = [
            [cell.text for cell in row.cells]
            for row in document.tables[0].rows
        ]

        self.assertEqual(paragraphs, EXPECTED_PARAGRAPHS)
        self.assertEqual(rows, EXPECTED_ROWS)
        with ZipFile(OUTPUT) as archive:
            self.assertIsNone(archive.testzip())


if __name__ == "__main__":
    unittest.main()
